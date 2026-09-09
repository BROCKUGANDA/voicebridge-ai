# VoiceBridge AI — 10-slide pitch deck for the AssemblyAI Voice Agent Hackathon.
# Brand: dark UI cyan/teal, Space Grotesk-ish (Calibri fallback for PPT).
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path

OUT = Path("download/VoiceBridge-AI-Pitch.pptx")
OUT.parent.mkdir(parents=True, exist_ok=True)

# Brand tokens
BG = RGBColor(0x0A, 0x0E, 0x17)
PANEL = RGBColor(0x0D, 0x11, 0x17)
INK = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x9A, 0xA4, 0xB2)
ACCENT = RGBColor(0x00, 0xF0, 0xFF)
TEAL = RGBColor(0x5C, 0xE0, 0xD8)
GOLD = RGBColor(0xFF, 0xB6, 0x27)
GREEN = RGBColor(0x51, 0xCF, 0x66)
RED = RGBColor(0xFF, 0x6B, 0x6B)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])

def fill_bg(slide, color=BG):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)

def textbox(slide, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return tb

def multi(slide, x, y, w, h, lines, size=16, color=MUTED, leading=1.15):
    """lines: list of (text, bold, color) or plain str"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in lines:
        if isinstance(item, str):
            text, bold, col = item, False, color
        else:
            text, bold, col = item
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        # Always set text via a run — empty strings have no runs[0]
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = col
        run.font.name = "Calibri"
        p.space_after = Pt(8)
    return tb

def accent_bar(slide, x=0.5, y=0.35, w=0.12, h=0.7):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = ACCENT
    sh.line.fill.background()

def pill(slide, x, y, w, h, text, fill=ACCENT, fg=BG, size=12):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = fg
    r.font.name = "Calibri"

def stat(slide, x, y, w, number, label, num_color=ACCENT):
    textbox(slide, x, y, w, 0.9, number, size=48, bold=True, color=num_color)
    textbox(slide, x, y + 0.85, w, 0.5, label, size=14, color=MUTED)

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

# ── Slide 1: Title ──────────────────────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s, 0.6, 1.8, 0.14, 1.6)
pill(s, 0.6, 1.0, 3.2, 0.36, "ASSEMBLYAI VOICE AGENT HACKATHON")
textbox(s, 0.9, 2.0, 11, 1.0, "VoiceBridge AI", size=54, bold=True, color=INK)
textbox(s, 0.9, 3.1, 11, 0.7, "Voice-first accessibility companion", size=28, color=ACCENT)
textbox(s, 0.9, 4.0, 10, 1.2, "Navigate a live web dashboard by voice.\nChange settings, read metrics, keep a full transcript —\nwith offline fallback when the cloud is down.", size=18, color=MUTED)
textbox(s, 0.9, 6.4, 10, 0.4, "Built on the AssemblyAI Voice Agent API  ·  Sep 2026", size=14, color=MUTED)
notes(s, "VoiceBridge AI. Voice-first accessibility companion built on AssemblyAI. 30 seconds: problem, live demo, why it wins.")

# ── Slide 2: Problem ────────────────────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s)
textbox(s, 0.9, 0.45, 11, 0.7, "Pixel-first UIs lock people out", size=36, bold=True)
multi(s, 0.9, 1.5, 5.8, 4.5, [
    ("Who hurts", True, ACCENT),
    ("Blind and low-vision users", False, INK),
    ("Motor-impaired users who cannot use a mouse", False, INK),
    ("Cognitively-diverse users overwhelmed by dense UIs", False, INK),
    ("", False, INK),
    ("Why current voice demos fail", True, GOLD),
    ("Chat wrappers that talk but do not operate a product", False, MUTED),
    ("Cloud-only — useless offline or during outages", False, MUTED),
    ("No tools: they cannot write prefs or read real metrics", False, MUTED),
], size=18)
# right stats
stat(s, 8.0, 1.8, 4, "1.3B", "people with significant disabilities (WHO)", ACCENT)
stat(s, 8.0, 3.5, 4, "WCAG", "legal pressure via EN 301 549 / ADA", TEAL)
stat(s, 8.0, 5.2, 4, "0", "acceptable excuses for unusable dashboards", GOLD)
notes(s, "Accessibility users are locked out of mouse-first dashboards. Existing voice demos are chatbots that cannot operate real product state, and they die offline.")

# ── Slide 3: Solution ───────────────────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s)
textbox(s, 0.9, 0.45, 12, 0.7, "Speak. The product acts. It talks back.", size=36, bold=True)
# three columns
cols = [
    (0.9, "1 · Listen", "PCM16 24 kHz mic\n→ AssemblyAI Voice Agent\nsingle-use token WS", ACCENT),
    (4.7, "2 · Act", "Agent calls dispatch_tool\n→ preferences, metrics,\nconversation, session", TEAL),
    (8.5, "3 · Confirm", "Short spoken reply\n+ live transcript\n+ screen-reader announce", GOLD),
]
for x, title, body, col in cols:
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.6), Inches(3.5), Inches(3.2))
    sh.fill.solid(); sh.fill.fore_color.rgb = PANEL
    sh.line.color.rgb = col
    textbox(s, x + 0.25, 1.85, 3, 0.5, title, size=20, bold=True, color=col)
    textbox(s, x + 0.25, 2.5, 3, 2, body, size=16, color=MUTED)
textbox(s, 0.9, 5.2, 11.5, 1.2,
        "Offline ladder when AssemblyAI is down:  on-device Vosk  →  browser Web Speech  →  typed commands.\nThe product never goes dark.",
        size=18, color=INK)
notes(s, "One WebSocket. One tool. Offline ladder. This is an operable product surface, not a chat toy.")

# ── Slide 4: Live demo screenshot ───────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s)
textbox(s, 0.9, 0.45, 12, 0.7, "Live product — not a mock", size=36, bold=True)
# embed frame if available
frame = Path("download/demo-frames/10-home-final.png")
if not frame.exists():
    frame = Path("download/demo-frames/03-home-ready.png")
if frame.exists():
    s.shapes.add_picture(str(frame), Inches(1.5), Inches(1.4), width=Inches(10.3))
else:
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.3), Inches(5))
    sh.fill.solid(); sh.fill.fore_color.rgb = PANEL; sh.line.fill.background()
    textbox(s, 2, 3.5, 9, 1, "[Dashboard screenshot]", size=24, color=MUTED, align=PP_ALIGN.CENTER)
textbox(s, 0.9, 6.7, 11, 0.4, "Power button  ·  engine health chip  ·  typed fallback  ·  8 panels  ·  5 settings tabs", size=14, color=MUTED)
notes(s, "Show the real UI. Power on, open metrics, change speed, read metrics. All clickable.")

# ── Slide 5: Why AssemblyAI ─────────────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s)
textbox(s, 0.9, 0.45, 12, 0.7, "Why AssemblyAI Voice Agent API", size=36, bold=True)
rows = [
    ("Full-duplex WS", "Listen + speak in one socket — interruption built in"),
    ("Stored agent_id", "Prompt, voice, keyterms, tools configured once via POST /v1/agents"),
    ("Single-use tokens", "API key never leaves the server; browser only gets a short-lived token"),
    ("Function tools", "dispatch_tool → tool.call / tool.result writes real product state"),
    ("Session lifecycle", "session.end on power-off + pagehide; auto-resume in the 30s window"),
]
y = 1.5
for label, desc in rows:
    pill(s, 0.9, y, 2.4, 0.38, label, fill=PANEL, fg=ACCENT, size=12)
    textbox(s, 3.5, y - 0.05, 9, 0.5, desc, size=16, color=INK)
    y += 0.85
textbox(s, 0.9, 6.3, 11, 0.6, "Live agent:  VoiceBridge Accessibility Agent  ·  voice: anna  ·  tool: dispatch_tool", size=14, color=MUTED)
notes(s, "We use the recommended stored-agent path, not just a token mint. Tools are real.")

# ── Slide 6: Differentiation ────────────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s)
textbox(s, 0.9, 0.45, 12, 0.7, "More than a voice wrapper", size=36, bold=True)
items = [
    ("Operates the product", "Opens panels, writes prefs, reads metrics — not just answers questions", ACCENT),
    ("Offline ladder", "Vosk WASM → Web Speech → typed input when the cloud fails", TEAL),
    ("Production security", "Signed sessions, CSRF, rate limits, AES-GCM, audit log, fail-closed secrets", GOLD),
    ("Accessibility-native", "Space / Alt+1–7, SR announcements, mute, transcript export, short replies", GREEN),
]
y = 1.5
for title, body, col in items:
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(y), Inches(11.5), Inches(1.05))
    sh.fill.solid(); sh.fill.fore_color.rgb = PANEL
    sh.line.color.rgb = col
    textbox(s, 1.2, y + 0.15, 10.5, 0.4, title, size=18, bold=True, color=col)
    textbox(s, 1.2, y + 0.55, 10.5, 0.4, body, size=14, color=MUTED)
    y += 1.25
notes(s, "Four differentiators judges care about: real actions, offline, security, a11y.")

# ── Slide 7: Market ─────────────────────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s)
textbox(s, 0.9, 0.45, 12, 0.7, "Who pays, and why now", size=36, bold=True)
stat(s, 0.9, 1.6, 3.5, "1.3B", "global disability market (WHO)", ACCENT)
stat(s, 4.5, 1.6, 3.5, "$26B+", "assistive tech trajectory", TEAL)
stat(s, 8.2, 1.6, 3.5, "Legal", "WCAG / EN 301 549 / ADA mandates", GOLD)
multi(s, 0.9, 3.5, 11, 3, [
    ("Primary buyers", True, ACCENT),
    ("Enterprise accessibility & digital inclusion teams", False, INK),
    ("Government / education under procurement a11y rules", False, INK),
    ("Product orgs shipping voice-first or hands-free UX", False, INK),
    ("", False, INK),
    ("Why now: regulation + LLM agents + browser audio matured together", True, TEAL),
], size=18)
notes(s, "Not 'everyone'. Accessibility teams with legal pressure. Voice agents finally work in the browser.")

# ── Slide 8: Business model ─────────────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s)
textbox(s, 0.9, 0.45, 12, 0.7, "Simple SaaS, AssemblyAI-metered cost", size=36, bold=True)
tiers = [
    (0.9, "Personal", "$12/mo", "Browser agent\nFallback engines\nTranscript export", ACCENT),
    (4.7, "Team", "$49/seat", "Shared prefs\nSession history\nAdmin audit", TEAL),
    (8.5, "Enterprise", "Custom", "SSO + private deploy\nTwilio phone\nSLA + support", GOLD),
]
for x, name, price, body, col in tiers:
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.5), Inches(3.5), Inches(4.2))
    sh.fill.solid(); sh.fill.fore_color.rgb = PANEL
    sh.line.color.rgb = col
    textbox(s, x + 0.3, 1.75, 3, 0.4, name, size=18, bold=True, color=col)
    textbox(s, x + 0.3, 2.3, 3, 0.7, price, size=36, bold=True, color=INK)
    textbox(s, x + 0.3, 3.2, 3, 2.2, body, size=16, color=MUTED)
textbox(s, 0.9, 6.1, 11.5, 0.6, "Unit economics: AssemblyAI Voice Agent minutes are the COGS. Offline fallback lowers billable cloud minutes.", size=14, color=MUTED)
notes(s, "Three tiers. AssemblyAI is metered cost. Fallback reduces cloud spend.")

# ── Slide 9: Build status ───────────────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s)
textbox(s, 0.9, 0.45, 12, 0.7, "Shipped — not vapor", size=36, bold=True)
proofs = [
    ("✓", "Stored AssemblyAI agent live (agent_id + dispatch_tool)"),
    ("✓", "Browser WS client: tokens, PCM16, tools, interrupt, resume"),
    ("✓", "Security stack: auth, CSRF, rate limit, AES-GCM, audit"),
    ("✓", "Typecheck clean  ·  lint clean  ·  8/8 unit tests  ·  prod build"),
    ("✓", "Playwright E2E: login, 8 panels, 5 tabs, Space, typed commands"),
    ("✓", "Frame-by-frame demo video + this deck"),
]
y = 1.5
for mark, text in proofs:
    textbox(s, 1.0, y, 0.5, 0.4, mark, size=20, bold=True, color=GREEN)
    textbox(s, 1.6, y, 10.5, 0.45, text, size=18, color=INK)
    y += 0.75
notes(s, "Real repo, real tests, real agent. E2E flyover captured.")

# ── Slide 10: Close ─────────────────────────────────────────────
s = blank(); fill_bg(s)
accent_bar(s, 0.6, 2.2, 0.14, 1.8)
textbox(s, 0.9, 2.0, 11, 0.8, "VoiceBridge AI", size=48, bold=True)
textbox(s, 0.9, 3.0, 11, 0.6, "Navigate. Configure. Understand — by voice.", size=24, color=ACCENT)
multi(s, 0.9, 4.0, 11, 2, [
    ("Try the live demo  ·  Read the repo  ·  Talk to the agent", True, INK),
    ("Roadmap: Twilio phone, multi-user auth, richer vision tools, enterprise SSO", False, MUTED),
    ("Thank you.", True, TEAL),
], size=18)
textbox(s, 0.9, 6.5, 11, 0.4, "lablab.ai  ·  AssemblyAI Voice Agent Hackathon  ·  Sep 2026", size=12, color=MUTED)
notes(s, "Close with the one-liner and the ask. Demo URL + repo + agent.")

prs.save(str(OUT))
print("WROTE", OUT, "slides", len(prs.slides.__iter__.__self__._sldIdLst))
